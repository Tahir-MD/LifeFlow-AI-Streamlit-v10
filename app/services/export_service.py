"""
Report downloads as Excel (.xlsx) and PDF.

Both take the same simple input: a title and a list of tables
    sheets = [{"title": "Requests", "headers": [...], "rows": [[...], ...]}, ...]
so every report (admin, hospital, donor) uses one code path. Needs `openpyxl` and `reportlab`
(both listed in requirements.txt).
"""
import io
from datetime import datetime
from typing import Any, Dict, List


def _cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if value is None:
        return ""
    return value


def build_xlsx(title: str, sheets: List[Dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    head_fill = PatternFill("solid", fgColor="7A1F2B")
    for sheet in sheets:
        ws = workbook.create_sheet(title=(sheet["title"] or "Sheet")[:31])
        ws.append([f"LifeFlow - {title}"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([f"Created {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC"])
        ws.append([])
        ws.append(sheet["headers"])
        header_row = ws.max_row
        for col in range(1, len(sheet["headers"]) + 1):
            cell = ws.cell(row=header_row, column=col)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet["rows"]:
            ws.append([_cell(v) for v in row])
        widths = [len(str(h)) for h in sheet["headers"]]
        for row in sheet["rows"]:
            for i, v in enumerate(row):
                widths[i] = max(widths[i], min(len(str(_cell(v))), 45))
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width + 3
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _safe(text: Any) -> str:
    """The built-in PDF fonts only know Latin letters; anything else is shown as '?' instead of crashing."""
    return str(_cell(text)).encode("latin-1", "replace").decode("latin-1")


def build_pdf(title: str, sheets: List[Dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title=f"LifeFlow - {title}")
    styles = getSampleStyleSheet()
    story = [Paragraph(_safe(f"LifeFlow - {title}"), styles["Title"]),
             Paragraph(f"Created {datetime.utcnow().strftime('%Y-%m-%d %H:%M')} UTC", styles["Normal"]),
             Spacer(1, 0.4 * cm)]
    usable = landscape(A4)[0] - 2.4 * cm
    for sheet in sheets:
        story.append(Paragraph(_safe(sheet["title"]), styles["Heading2"]))
        headers = [_safe(h) for h in sheet["headers"]]
        if not sheet["rows"]:
            story.append(Paragraph("No records.", styles["Normal"]))
            story.append(Spacer(1, 0.3 * cm))
            continue
        small = styles["BodyText"].clone("small", fontSize=7.5, leading=9)
        data = [[Paragraph(f"<b>{h}</b>", small) for h in headers]]
        for row in sheet["rows"]:
            data.append([Paragraph(_safe(v).replace("&", "&amp;").replace("<", "&lt;"), small) for v in row])
        table = Table(data, colWidths=[usable / len(headers)] * len(headers), repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3E3E1")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCBBB8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FBF6F5")]),
        ]))
        story.append(table)
        story.append(Spacer(1, 0.5 * cm))
    doc.build(story)
    return buffer.getvalue()
